"""Bounded document extraction in a cancellable helper process."""

import base64
import io
import json
import sys
import zipfile


def main():
    request = json.load(sys.stdin)
    data = base64.b64decode(request["data"], validate=True)
    media = request["mediaType"]
    if media == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted or len(reader.pages) > 128:
            raise ValueError("encrypted or excessive PDF")
        text = ""
        for page in reader.pages:
            text += (page.extract_text() or "") + "\n"
            if len(text) > 250_000:
                raise ValueError("extracted PDF exceeds limit")
    elif (
        media
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        from docx import Document

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 2048
                or sum(item.file_size for item in entries) > 32 * 1024 * 1024
            ):
                raise ValueError("DOCX archive exceeds expansion limit")
        document = Document(io.BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        text += "\n" + "\n".join(
            "\t".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
    else:
        raise ValueError("unsupported document media type")
    if len(text) > 250_000:
        raise ValueError("extracted document exceeds limit")
    print(json.dumps({"text": text}, ensure_ascii=True))


if __name__ == "__main__":
    main()
